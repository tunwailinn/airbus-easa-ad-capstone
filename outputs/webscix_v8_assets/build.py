from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
D=Path('/Users/tunwailin/Documents/Capstone/outputs')
A=D/'webscix_v8_assets'
fontdir=Path('/System/Library/Fonts/Supplemental')
def font(s=40,b=False):return ImageFont.truetype(str(fontdir/('Times New Roman Bold.ttf' if b else 'Times New Roman.ttf')),s)
ink='#18232B'; blue='#365E79'; grey='#D7DFE4'
def canvas(h):
 im=Image.new('RGB',(1440,h),'white');return im,ImageDraw.Draw(im)
def txt(x,y,t,s=40,b=False,c=ink):dr.text((x,y),t,font=font(s,b),fill=c)
def center(y,t,s=40,b=False):
 w=dr.textlength(t,font=font(s,b));txt((1440-w)/2,y,t,s,b)
im,dr=canvas(1170)
center(10,'EASA Airbus corpus',46,True)
center(70,'1,809 PDFs covering 1,808 distinct base AD numbers',38)
center(120,'Operational view: 1,786 PDFs · 6,002 verified pages',38)
for i,(title,l1,l2) in enumerate([
 ('Identity and source preparation','Parser v2.1.6 · lifecycle records · page-text v1.1','12,634 section chunks · source-fidelity checks'),
 ('Evidence retrieval','Known-document routing or identifier-free discovery','BM25 + Qwen3 fusion → rerank 20 → top 5'),
 ('Evidence-constrained answers','DeepSeek V4 Pro · structured-response validation','Local citation checks · explicit abstention')]):
 y=215+i*255
 dr.rounded_rectangle((25,y,1415,y+195),radius=12,outline=grey,width=3,fill='#F8FAFB')
 dr.ellipse((60,y+64,127,y+131),fill=blue)
 txt(79,y+70,chr(65+i),43,True,'white')
 txt(164,y+22,title,45,True)
 txt(164,y+83,l1,38)
 txt(164,y+133,l2,38)
 if i<2:
  dr.line((720,y+204,720,y+239),fill=blue,width=4)
  dr.polygon([(709,y+230),(731,y+230),(720,y+244)],fill=blue)
center(1015,'Evaluation after the research freeze',40,True)
center(1068,'Primary benchmark: 40 questions · Unseen study: 5 PDFs / 15 questions',36)
center(1115,'Serving optimization follows the research freeze',36)
im.save(A/'architecture.png',dpi=(300,300))
# Independent panels preserve distinct evaluation populations.
im,dr=canvas(850)
for y,title,sub,rows in [
 (10,'(a) Historical benchmark — QA-v2','44 answerable questions; semantic accuracy not evaluated', [('E0 · Recall@5',0,'0/44 · 0.00%'),('E4 · Recall@5',18/44*100,'18/44 · 40.91%')]),
 (430,'(b) Frozen E5-D final benchmark','36 answerable retrieval questions / 40 total questions',[('Recall@5',35/36*100,'35/36 · 97.22%'),('Semantic accuracy',95,'38/40 · 95.00%')])]:
 txt(25,y,title,44,True);txt(25,y+59,sub,37)
 left,right=390,1110
 for tick in [0,25,50,75,100]:
  x=left+(right-left)*tick/100
  dr.line((x,y+125,x,y+285),fill='#E3E7EB',width=2)
  txt(x-18,y+300,str(tick),32)
 for j,(label,val,value) in enumerate(rows):
  yy=y+148+j*90
  txt(25,yy-7,label,38)
  if val:dr.rectangle((left,yy,left+(right-left)*val/100,yy+48),fill=blue if j==0 else '#819BAC')
  else:dr.ellipse((left-5,yy+19,left+5,yy+29),fill=blue)
  txt(1138,yy+3,value,33,True)
 txt(657,y+353,'Score (%)',35)
im.save(A/'benchmark.png',dpi=(300,300))
d=Document(D/'WebSciX2026_Full_Paper_Polished_v7.docx');ps=d.paragraphs
# Replace first inline drawing, preserving the existing paragraph.
p=ps[15];p.clear();s=p.add_run().add_picture(str(A/'architecture.png'),width=Inches(4.8));s._inline.docPr.set('descr','Three-layer architecture from corpus preparation through retrieval and evidence-constrained answers. Primary and unseen evaluation remain separate.')
# Existing figures following new figure shift by one.
for p in d.paragraphs:
 for r in p.runs:
  updated=r.text.replace('Fig. 5.','Fig. 6.').replace('Fig. 4.','Fig. 5.').replace('Table 2','Fig. 4')
  if updated!=r.text:r.text=updated
# Place new chart before the caption, remove former table.
p=ps[48];p.text='Fig. 4. Historical E0/E4 retrieval results and the separate E5-D final benchmark. Different question sets preclude a controlled cross-set ablation.'
caption=p.text;p.clear();p.add_run('Fig. 4.').bold=True;p.add_run(caption[len('Fig. 4.'):])
for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9)
new=p.insert_paragraph_before();new.alignment=WD_ALIGN_PARAGRAPH.CENTER;new.paragraph_format.keep_with_next=True
s=new.add_run().add_picture(str(A/'benchmark.png'),width=Inches(4.8));s._inline.docPr.set('descr','Historical QA-v2 Recall@5: E0 0 of 44, E4 18 of 44. Separate E5 final Recall@5: 35 of 36; semantic accuracy: 38 of 40. Scores are not a controlled comparison across datasets.')
t=d.tables[1];t._element.getparent().remove(t._element)
ps[49].text='Configurations: E0 uses flat chunks (≤350 units), MiniLM-L6-v2 dense retrieval and no reranker. E4 uses section-aware chunks (≤450 units), BM25 + MiniLM-L6-v2 fusion (RRF) and a MiniLM cross-encoder. E5-D uses section-aware chunks (≤450 units), BM25 + Qwen3-0.6B fusion (RRF) and a Qwen3 reranker. Chunk units are deterministic whitespace units, not model tokens.'
ps[49].paragraph_format.keep_with_next=False
ps[49].paragraph_format.keep_together=True
for r in ps[49].runs:r.font.name='Times New Roman';r.font.size=Pt(9)
# Caption now follows chart; no need to bind it to long configuration note.
p.paragraph_format.keep_with_next=False
out=D/'WebSciX2026_Full_Paper_Polished_v8.docx';d.save(out);print(out)
