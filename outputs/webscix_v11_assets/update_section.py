from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
R=Path('/Users/tunwailin/Documents/Capstone/outputs')
d=Document(R/'WebSciX2026_Full_Paper_Polished_v10.docx')
ps=d.paragraphs
start=next(i for i,p in enumerate(ps) if p.text.startswith('3.3 '))
end=next(i for i,p in enumerate(ps) if p.text.startswith('3.4 '))
base=deepcopy(ps[start+1]._p.pPr)
mathp=deepcopy(next(p._p for p in ps[start:end] if p._p.xpath('.//m:oMath')))
blocks=[
'The retrieval pipeline uses different strategies depending on whether the user asks about a specific Airworthiness Directive or performs corpus-wide discovery. For a **known-document query**, the user explicitly provides an AD identifier. The system deterministically recognizes that identifier and restricts retrieval to passages from the corresponding directive. This routing step does not require an LLM. For a **discovery query**, where no target AD number is provided, the system must search the complete operational corpus to identify both the relevant directive and the most relevant passages.',
'For identifier-free discovery, the **E2-C** candidate generator combines two complementary retrieval methods: **BM25 sparse retrieval** and **Qwen3 dense retrieval**. BM25 ranks documents according to lexical overlap between the user question and the indexed text. This is useful for exact engineering expressions such as part numbers, modification numbers, aircraft identifiers, and regulatory terminology. In parallel, the question is encoded using **Qwen3-Embedding-0.6B**, and the resulting vector is compared with precomputed chunk embeddings using cosine similarity. The dense retrieval branch is intended to capture semantic similarity when the user’s wording differs from the terminology used in the AD.',
'Because BM25 scores and cosine-similarity scores are produced on different numerical scales, they are not combined directly. Instead, their ranked results are merged using **Reciprocal Rank Fusion (RRF)** [6]. For a candidate document d, the fused score is',
None,
'where M = {BM25, Dense}, rₘ(d) is the rank assigned to document d by retrieval method m, and k = 60 is the fixed RRF constant. A document receives a higher fused score when it appears near the top of one or both ranked lists. Using rank positions rather than raw scores allows lexical and dense retrieval to be combined without requiring their scoring scales to be directly comparable.',
'The highest-ranked ADs are then shortlisted, and retrieval is performed again at the **passage level within those documents**. BM25 and dense passage rankings are fused using the same RRF procedure. This second fusion stage produces a broad set of candidate passages that benefits from both exact lexical matching and semantic similarity. Dense embeddings are generated in advance for the frozen section-aware chunk collection, and the stored vector order and normalization are verified so that each embedding remains correctly associated with its original source chunk.',
'The hybrid E2-C stage is designed to maximize the quality of the candidate pool rather than to define the final evidence order. The **20 highest-ranked candidate passages** are therefore passed to **E2-D**, which uses **Qwen3-Reranker-0.6B** for a second-stage relevance assessment. The reranker evaluates each candidate together with the user question and assigns a new relevance score. The candidates are then reordered according to these scores, and only the **five highest-ranked passages** are retained as the final evidence set supplied to Layer C for answer generation. In this design, **E2-C identifies plausible candidate evidence, while E2-D determines which five passages are most relevant to the question**.',
'The reranker operates under a **fixed engineering-focused instruction** that is not modified after the evaluation configuration is frozen. For reproducibility, the exact model revisions are also fixed: dense retrieval uses **Qwen/Qwen3-Embedding-0.6B at revision 97b0c61**, while reranking uses **Qwen/Qwen3-Reranker-0.6B at revision e61197e**. Pinning these revisions prevents later changes to the public model repositories from silently altering retrieval behavior.',
'Each of the five selected evidence passages retains its original provenance metadata, including the **chunk identifier, AD number, page number, section, source PDF, and final rank**. Therefore, the evidence passed to Layer C is fully traceable to the original regulatory source rather than being treated as anonymous text.'
]
for p in ps[start+1:end]:p._p.getparent().remove(p._p)
anchor=ps[start]._p
for text in blocks:
 if text is None:
  node=mathp
 else:
  node=OxmlElement('w:p');node.append(deepcopy(base));p=Paragraph(node,d._body)
  p.paragraph_format.keep_together=False
  p.paragraph_format.keep_with_next=False
  for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
   if not part:continue
   r=p.add_run(part);r.bold=bool(i%2);r.font.name='Times New Roman';r.font.size=Pt(10)
 anchor.addnext(node);anchor=node
# Avoid separating the equation lead-in from its display.
p=next(p for p in d.paragraphs if p.text.startswith('Because BM25 scores'))
p.paragraph_format.keep_with_next=True
out=R/'WebSciX2026_Full_Paper_Polished_v11.docx'
d.save(out)
print(out)
