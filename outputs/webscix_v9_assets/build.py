from pathlib import Path
import json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
D=Path('/Users/tunwailin/Documents/Capstone/outputs');A=D/'webscix_v9_assets'
F=Path('/System/Library/Fonts/Supplemental')
def font(s=38,b=False):return ImageFont.truetype(str(F/('Arial Bold.ttf' if b else 'Arial.ttf')),s)
ink='#202A32';blue='#315E7A';teal='#367E77';gray='#849AA9'
def canvas(h):
 global im,dr
 im=Image.new('RGB',(1440,h),'white');dr=ImageDraw.Draw(im)
def txt(x,y,t,s=38,b=False,c=ink):dr.text((x,y),t,font=font(s,b),fill=c)
def center(y,t,s=38,b=False):txt((1440-dr.textlength(t,font=font(s,b)))/2,y,t,s,b)
def save(name):im.save(A/name,dpi=(300,300))
canvas(1090)
center(10,'EASA Airbus corpus',44,True)
center(68,'1,809 PDFs · 1,808 distinct base AD numbers',36)
center(118,'Operational view: 1,786 PDFs · 6,002 verified pages',36)
for i,(title,l1,l2) in enumerate([
 ('Identity and source preparation','Parser v2.1.6 · lifecycle records · page-text v1.1','12,634 section chunks · source-fidelity checks'),
 ('Evidence retrieval','Known-document routing or identifier-free discovery','BM25 + Qwen3 fusion → rerank 20 → top 5'),
 ('Evidence-constrained answers','DeepSeek V4 Pro · structured-response validation','Local citation checks · explicit abstention')]):
 y=200+i*237
 dr.rounded_rectangle((25,y,1415,y+182),radius=12,outline='#CBD5DC',width=3,fill='#F7F9FA')
 dr.ellipse((57,y+55,127,y+125),fill=blue);txt(78,y+65,chr(65+i),38,True,'white')
 txt(164,y+20,title,42,True);txt(164,y+79,l1,35);txt(164,y+127,l2,35)
 if i<2:
  dr.line((720,y+193,720,y+222),fill=blue,width=4);dr.polygon([(709,y+216),(731,y+216),(720,y+230)],fill=blue)
center(930,'Evaluation boundary',38,True)
center(987,'Primary: 40 questions · Unseen: 5 PDFs / 15 questions',35)
center(1040,'Serving optimization follows the research freeze',35)
save('fig1.png')
def bars(name,labels,values,texts,h=650,xmax=100,unit='Score (%)',colors=None):
 canvas(h);left,right=580,1220;top=35;bottom=h-145
 step=(bottom-top)/len(labels)
 ticks=[0,25,50,75,100] if xmax==100 else [0,10,20,30]
 for t in ticks:
  x=left+(right-left)*t/xmax;dr.line((x,top-5,x,bottom+3),fill='#E3E7EB',width=2);txt(x-17,bottom+19,str(t),34)
 for j,(lab,v,val) in enumerate(zip(labels,values,texts)):
  y=top+j*step+step*.18;tw=dr.textlength(lab,font=font(37));txt(left-30-tw,y,lab,37)
  end=left+(right-left)*v/xmax;dr.rectangle((left,y,end,y+step*.52),fill=(colors or [blue]*len(labels))[j]);txt(end+13,y,val,34)
 txt(815,h-58,unit,35);save(name)
bars('fig2.png',['Stable metadata F1','Applicability model F1','Reference number F1','Superseded AD F1','Source containment'],[98.31,92.22,90,66.67,100],['98.31','92.22','90.00','66.67','100.00'],h=615,colors=[blue]*4+[teal])
reports=[json.loads((D.parent/f'data_processed/evaluations/e5/e5{x}_development_evaluation.json').read_text()) for x in 'abcd']
assert len({r['benchmark_sha256'] for r in reports})==1
assert all(r['question_accounting']['answerable_retrieval']==54 for r in reports)
canvas(1090)
for y,title,sub,labels,vals,counts in [
 (5,'(a) Historical retrieval benchmark','QA-v2 · 44 answerable questions',['E0','E4'],[0,18/44*100],['0/44','18/44']),
 (430,'(b) E5 development ablation','Same 54 answerable questions for E5-A through E5-D',[r['experiment'] for r in reports],[r['overall']['recall_at_5']*100 for r in reports],[f"{round(r['overall']['recall_at_5']*54)}/54" for r in reports])]:
 txt(25,y,title,42,True);txt(25,y+60,sub,35)
 left,right=275,1130;top=y+139;bottom=top+len(vals)*91-15
 for t in [0,25,50,75,100]:
  x=left+(right-left)*t/100;dr.line((x,top-8,x,bottom),fill='#E3E7EB',width=2);txt(x-18,bottom+18,str(t),32)
 for j,(lab,v,count) in enumerate(zip(labels,vals,counts)):
  yy=top+j*91;txt(30,yy+2,lab,38,lab=='E5-D')
  if v:dr.rectangle((left,yy,left+(right-left)*v/100,yy+49),fill=blue if lab=='E5-D' else gray)
  else:dr.ellipse((left-5,yy+20,left+5,yy+30),fill=gray)
  txt(1153,yy+4,f'{v:.2f}%',35,True);txt(1310,yy+7,count,30)
 txt(620,bottom+64,'Recall@5 (%)',34)
save('fig3.png')
bars('fig4.png',['Recall@1','Recall@3','Recall@5 overall','Known-document R@5','Discovery R@5','Semantic accuracy'],[83.33,97.22,97.22,100,91.67,95],['83.33','97.22','97.22','100.00','91.67','95.00'],h=675,colors=[blue]*5+[teal])
canvas(795)
# Grouped final/unseen comparison, same explicit denominators as the source.
dr.rectangle((315,20,356,45),fill=blue);txt(370,10,'Frozen final',35)
dr.rectangle((695,20,736,45),fill=teal);txt(750,10,'Unseen after ingestion',35)
left,right,top,bottom=140,1390,155,590
for t in [0,25,50,75,100]:
 y=bottom-(bottom-top)*t/100;dr.line((left,y,right,y),fill='#E3E7EB',width=2);txt(62,y-18,str(t),34)
for j,(label,a,b,ca,cb) in enumerate([('Retrieval\nRecall@5',97.22,100,'35/36','14/14'),('Semantic\naccuracy*',95,92.86,'38/40','13/14'),('Strict\nend-to-end',95,86.67,'38/40','13/15')]):
 x=245+j*410
 for xx,v,count,c in [(x,a,ca,blue),(x+150,b,cb,teal)]:
  y=bottom-(bottom-top)*v/100;dr.rectangle((xx,y,xx+137,bottom),fill=c);txt(xx+5,y-43,f'{v:.2f}',33);txt(xx+8,bottom-60,count,34,True,'white')
 for k,line in enumerate(label.split('\n')):txt(x+130-dr.textlength(line,font=font(35))/2,620+k*45,line,35)
txt(42,95,'Score (%)',32);save('fig5.png')
bars('fig6.png',['Legacy subprocess','Persistent warm serving'],[26.87,6.11],['26.87 s','6.11 s'],h=450,xmax=30,unit='Median retrieval latency (s)',colors=[gray,teal])
# Use native paragraphs for scholarly captions and narrative; only figure labels use Arial.
d=Document(D/'WebSciX2026_Full_Paper_Polished_v8.docx');ps=d.paragraphs
picps=[p for p in ps if p._p.xpath('.//w:drawing')]
for i,p in enumerate(picps,1):
 p.clear();sh=p.add_run().add_picture(str(A/f'fig{i}.png'),width=Inches(4.8));sh._inline.docPr.set('descr',[
 'Three-layer architecture: corpus preparation, evidence retrieval and evidence-constrained answers.',
 'Clean extraction F1 scores and audited source containment.',
 'Recall at five: E0 0/44 and E4 18/44 on QA-v2. Separate E5 development set: E5-A 48/54, E5-B 51/54, E5-C 50/54, E5-D 52/54.',
 'Frozen E5 final retrieval metrics and 38/40 semantic accuracy.',
 'Final versus unseen results with sample sizes; semantic accuracy excludes provider failures, strict success includes them.',
 'Median retrieval latency: legacy subprocess 26.87 seconds and persistent warm serving 6.11 seconds.'][i-1])
def settext(p,t,size=None):
 p.text=t
 for r in p.runs:
  r.font.name='Times New Roman'
  if size:r.font.size=Pt(size)
def caption(p,t):
 p.clear();lead,rest=t.split('. ',1);p.add_run(lead+'.').bold=True;p.add_run(' '+rest)
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9)
settext(ps[44],'5.2 Retrieval Development and Frozen Final QA')
settext(ps[45],'The retrieval experiments progressed from E0 and E4 to four E5 development stages (Fig. 3). E0 used flat chunks (≤350 whitespace units), MiniLM-L6-v2 dense retrieval and no reranker. E4 used section-aware chunks (≤450 units), BM25–MiniLM fusion and a MiniLM cross-encoder. E5-A introduced deterministic AD routing, BM25 and section preferences. E5-B preserved known-document retrieval while adding two-stage sparse discovery and evidence assembly. E5-C added Qwen3-Embedding-0.6B document and passage fusion for discovery. E5-D reranked the fixed E5-C top-20 candidates using Qwen3-Reranker-0.6B. All E5 stages reused the E4 section chunks with an evidence depth of five; chunk units are whitespace units, not model tokens.')
caption(ps[47],'Fig. 3. Historical E0/E4 results and E5-A–D development ablations. Recall@5 uses 44 and 54 answerable questions, respectively. Comparisons across panels are not controlled ablations; the E5 final test is reported separately in Fig. 4.')
caption(ps[49],'Fig. 4. Frozen E5-D final results. Overall retrieval uses 36 answerable questions; known-document and discovery subsets contain 24 and 12 questions. Semantic accuracy uses all 40 questions.')
settext(ps[50],'On the common E5 development set, Recall@5 was 48/54 (88.89%) for E5-A, 51/54 (94.44%) for E5-B, 50/54 (92.59%) for E5-C and 52/54 (96.30%) for E5-D. The dense signal did not improve top-five recall over E5-B, although candidate source-and-page recall at depth 20 rose from 52/54 to 53/54. Reranking then improved the ordering of that fixed candidate pool. E5-D achieved MRR@5 of 0.8633 and nDCG@5 of 0.8884, with known-document Recall@5 of 36/36 and discovery Recall@5 of 16/18. E5-D was selected and frozen before the final test.')
ps[50].paragraph_format.keep_together=False
settext(ps[51],'The separate final benchmark produced 35/36 Recall@5 (97.22%) on answerable retrieval questions and 38/40 human semantic passes (95.0%) end to end (Fig. 4). Known-document retrieval remained perfect at 24/24 (100%), while identifier-free discovery achieved 11/12 Recall@5 (91.67%). These final-test results are distinct from both the E5 development ablations and the historical QA-v2 experiment.')
ps[47]._p.addnext(ps[50]._p)
for r in ps[50].runs:r.font.size=Pt(10)
ps[50].paragraph_format.first_line_indent=ps[45].paragraph_format.first_line_indent
out=D/'WebSciX2026_Full_Paper_Polished_v9.docx';d.save(out)
(A/'verified_metrics.json').write_text(json.dumps([{k:r[k] for k in ['experiment','benchmark_sha256','question_accounting','overall','configuration']} for r in reports],indent=2))
print(out)
